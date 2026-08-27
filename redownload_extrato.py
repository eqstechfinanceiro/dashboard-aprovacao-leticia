import requests
import json
import psycopg2
import openpyxl
import io
import warnings
from datetime import datetime, timedelta
warnings.filterwarnings('ignore')

API_URL = 'https://api.vexpenses.com'
DB_CONN = "postgresql://neondb_owner:npg_iItZN95svyEG@ep-restless-voice-amrrz188-pooler.c-5.us-east-1.aws.neon.tech/neondb?sslmode=require"

LARAVEL_TOKEN = "eyJpdiI6IlBXOC9qcXpaQmlQTDBIUzJPWE95REE9PSIsInZhbHVlIjoiVUM5MTArSGQ1UXd5Y0xrbjQ2U3RhV2dONmtEZ2NzTVE1dmF2TEdMWkx5YitSMWszWkNGdk9oTHVnNW1xMkZhMVRHVTRjU2R5blBtUjVTNE1tK0JQTXByd2FIYmRCZmR3cS9pN01SRWVORjVldlhSUFVZNHpyZDZGLzJVbGk1QjBSano2anVqa0xzY1JiMVBCUHZKaHlJOTR5ek5FSFRRQkQ4NUJiWW5qRGxVSE4vb3dWN2pKaW9qbkJyM3NkY1VOZS9WV0d2TnFBQ2o3QlBudENkUUMzVy94SjcxWTFwVlFHdUIwNm5aZHZJWTdYWi9IYWU4cWxlZWhOT21HUTNGL0JVM2EyTDNod1dyN1pwd1dyUmVJc09SV01VNldEazBRellHOFdXREdvcDIvMHNRbDhMOEJWbzAzVWpLLzkxaXQiLCJtYWMiOiI3ODRhNmQwZmMzMjA3ZDkxOTM3N2NmOWRhNThlOTJjNmNhMzY4MjU4OWU4YzU3YjkwMTAzZmQ1Mjk0OWVjNWExIiwidGFnIjoiIn0%3D"
LARAVEL_SESSION = "m25jkuQuhvM3PsG8ENnYTauKMd8kbHmDIeyLSRqv"
XSRF_TOKEN = "eyJpdiI6IkFlYm82amw0Q3p0T2l5QnMzV0tRU1E9PSIsInZhbHVlIjoiZ09jLzYwWXdjaFdreXdjQ3BERXFvVTl2azRldFJaaUxvOFNHbENBYkVwbUpYYk05ZzFlc3FGRDAzeHphTEpSL3ZvdlFwRG83ajBKZ2NyWE44cW1xWFJaWXdudUJ1ZG90VTBmVWtWbVlrRGpCNXZpakxoYVBENXRYRU5MeC9rMzkiLCJtYWMiOiI2ZDRjMWJhYTIwNjM2ODBmYzRkNWQ3YWUwYTU2OTMxMzEyZmY2ZjE3Yjk4ZmY5MzUzNGI2ZDhkMTZkNTYzMGE3IiwidGFnIjoiIn0%3D"

COOKIE = f"laravel_token={LARAVEL_TOKEN}; laravel_session={LARAVEL_SESSION}; XSRF-TOKEN={XSRF_TOKEN}; language=pt-BR"

HEADERS = {
    'Cookie': COOKIE,
    'Accept': 'application/json',
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
    'Origin': 'https://app.vexpenses.com',
    'Referer': 'https://app.vexpenses.com',
}

# Generate 15-day chunks from 2025-01-01 to 2026-08-25
chunks = []
start = datetime(2025, 1, 1)
end = datetime(2026, 8, 25)
current = start
while current <= end:
    chunk_end = current + timedelta(days=14)
    if chunk_end > end:
        chunk_end = end
    chunks.append((current.strftime('%Y-%m-%d'), chunk_end.strftime('%Y-%m-%d')))
    current = chunk_end + timedelta(days=1)

print(f"Total chunks: {len(chunks)}")

conn = psycopg2.connect(DB_CONN)
cur = conn.cursor()

total_inserted = 0
total_skipped = 0

for i, (chunk_start, chunk_end) in enumerate(chunks):
    print(f"\nChunk {i+1}/{len(chunks)}: {chunk_start} to {chunk_end}")
    
    try:
        # Step 1: Get S3 URL
        url = f"{API_URL}/v3/pay/statement/excel-all?start_date={chunk_start}&end_date={chunk_end}"
        resp = requests.get(url, headers=HEADERS, timeout=90)
        
        if resp.status_code != 200:
            print(f"  API returned {resp.status_code}, skipping")
            continue
        
        data = resp.json()
        if not data.get('success'):
            print(f"  API error, skipping")
            continue
        
        s3_url = data.get('data', {}).get('url')
        if not s3_url:
            print(f"  No S3 URL, skipping")
            continue
        
        # Step 2: Download XLSX
        xlsx_resp = requests.get(s3_url, timeout=180)
        if xlsx_resp.status_code != 200:
            print(f"  XLSX download failed {xlsx_resp.status_code}")
            continue
        
        # Step 3: Parse XLSX
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        
        rows_iter = ws.iter_rows(min_row=1, values_only=True)
        header = next(rows_iter)
        header_map = {str(v).strip(): idx for idx, v in enumerate(header) if v}
        
        col_map = {
            'Data': 'data', 'Hora': 'hora', 'Código de Transação': 'codigo_transacao',
            'Número do Cartão': 'numero_cartao', 'Grupo': 'grupo', 'Usuário': 'usuario',
            'Tipo': 'tipo', 'Descrição': 'descricao', 'Valor': 'valor',
            'Status': 'status', 'ID da Despesa': 'id_despesa',
            'ID do Relatório': 'id_relatorio', 'Tipo de Despesa': 'tipo_despesa',
            'Centro de Custo': 'centro_custo', 'Projeto': 'projeto',
            'Percentual de projeto': 'percentual_projeto',
        }
        
        chunk_inserted = 0
        chunk_skipped = 0
        
        for row in rows_iter:
            if row is None:
                continue
            
            transformed = {}
            for xlsx_col, db_col in col_map.items():
                idx = header_map.get(xlsx_col)
                if idx is not None and idx < len(row) and row[idx] is not None:
                    transformed[db_col] = row[idx]
            
            # Convert date
            data_val = transformed.get('data')
            if isinstance(data_val, datetime):
                data_val = data_val.strftime('%Y-%m-%d')
            elif isinstance(data_val, (int, float)):
                data_val = (datetime(1899, 12, 30) + timedelta(days=data_val)).strftime('%Y-%m-%d')
            elif isinstance(data_val, str) and data_val.isdigit():
                serial = int(data_val)
                data_val = (datetime(1899, 12, 30) + timedelta(days=serial)).strftime('%Y-%m-%d')
            
            # Convert valor
            valor = transformed.get('valor')
            if isinstance(valor, str):
                valor = float(valor.replace('.', '').replace(',', '.'))
            
            is_snapshot = not transformed.get('tipo') or transformed.get('hora') == '-'
            
            # Check if this entry already exists (by codigo_transacao + data + valor)
            cod = transformed.get('codigo_transacao', '')
            if cod:
                cur.execute("""
                    SELECT 1 FROM extrato_movimentacao
                    WHERE codigo_transacao = %s AND data = %s AND valor = %s
                    LIMIT 1
                """, (cod, data_val, valor))
                if cur.fetchone():
                    chunk_skipped += 1
                    continue
            
            # Insert new entry
            cur.execute("""
                INSERT INTO extrato_movimentacao
                (data, hora, codigo_transacao, numero_cartao, grupo, usuario, tipo,
                 descricao, valor, status, id_despesa, id_relatorio, tipo_despesa,
                 centro_custo, projeto, percentual_projeto, is_snapshot)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """, (
                data_val, transformed.get('hora'), transformed.get('codigo_transacao'),
                transformed.get('numero_cartao'), transformed.get('grupo'),
                transformed.get('usuario'), transformed.get('tipo'),
                transformed.get('descricao'), valor, transformed.get('status'),
                transformed.get('id_despesa'), transformed.get('id_relatorio'),
                transformed.get('tipo_despesa'), transformed.get('centro_custo'),
                transformed.get('projeto'), transformed.get('percentual_projeto'),
                is_snapshot,
            ))
            chunk_inserted += 1
        
        conn.commit()
        total_inserted += chunk_inserted
        total_skipped += chunk_skipped
        print(f"  Inserted: {chunk_inserted}, Skipped (existing): {chunk_skipped}")
        
    except Exception as e:
        print(f"  Error: {e}")
        conn.rollback()

cur.close()
conn.close()

print(f"\n=== DONE ===")
print(f"Total inserted: {total_inserted}")
print(f"Total skipped (already existed): {total_skipped}")
