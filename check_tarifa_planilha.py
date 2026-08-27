import openpyxl
import warnings
warnings.filterwarnings('ignore')

PLANILHA = 'CONTROLE - VEXPENSES - AGOSTO 2026.xlsx'

cpfs = {
    '02027745203': 'ABNER',
    '08247635992': 'JORGE LUIZ',
    '11178519740': 'GUILHERME MOTTA',
    '00041171071': 'JAFER',
    '72756284220': 'MARCO AURELIO ANDRADE',
    '02299450076': 'ALEKSANDER',
    '47514649816': 'LUCAS TAVARES',
}

print("Lendo planilha...")
wb = openpyxl.load_workbook(PLANILHA, read_only=True, data_only=True)

ws = wb['EXTRATO']
print(f"EXTRATO: {ws.max_row} rows")

rows_iter = ws.iter_rows(min_row=8, values_only=True)
header_row = next(rows_iter)
headers = {str(v).strip().lower(): i for i, v in enumerate(header_row) if v}

tipo_idx = headers.get('tipo', 9)
valor_idx = headers.get('valor', 11)
cpf_idx = headers.get('cpf', 12)
data_idx = headers.get('data', 3)
desc_idx = headers.get('descrição', 10)
cod_idx = headers.get('código de transação', 5)

# Collect TARIFA entries for each CPF
entries = {cpf: [] for cpf in cpfs}
for row in rows_iter:
    if row is None:
        continue
    cpf_val = str(row[cpf_idx] or '').strip() if cpf_idx < len(row) else ''
    if cpf_val not in cpfs:
        continue
    tipo_val = str(row[tipo_idx] or '').strip() if tipo_idx < len(row) else ''
    if tipo_val != 'TARIFA':
        continue
    valor_val = float(row[valor_idx] or 0) if valor_idx < len(row) else 0.0
    data_val = str(row[data_idx] or '') if data_idx < len(row) else ''
    desc_val = str(row[desc_idx] or '') if desc_idx < len(row) else ''
    cod_val = str(row[cod_idx] or '') if cod_idx < len(row) else ''
    entries[cpf_val].append((data_val, valor_val, desc_val, cod_val))

for cpf, name in cpfs.items():
    ents = entries.get(cpf, [])
    total = sum(v for _, v, _, _ in ents)
    print(f"\n{name} ({cpf}): {len(ents)} TARIFA entries, sum={total}, abs={abs(total)}")
    for d, v, desc, cod in sorted(ents):
        print(f"  {d} | {v:>10.2f} | {desc[:60]} | {cod}")

# PAINEL
print("\n\nPAINEL tarifa values:")
ws_p = wb['PAINEL']
for row in ws_p.iter_rows(min_row=12, values_only=True):
    if row is None or row[2] is None:
        continue
    cpf = str(row[2]).strip()
    if cpf in cpfs:
        nome = str(row[1]).strip() if row[1] else ''
        tarifa = float(row[15]) if row[15] else 0.0
        carga = float(row[13]) if row[13] else 0.0
        print(f"  {cpf} {nome}: tarifa={tarifa}, carga={carga}")
