import openpyxl
import warnings
warnings.filterwarnings('ignore')

PLANILHA = 'CONTROLE - VEXPENSES - AGOSTO 2026.xlsx'

wb = openpyxl.load_workbook(PLANILHA, read_only=True, data_only=True)
ws = wb['EXTRATO']

# Read header to find column indices
rows_iter = ws.iter_rows(min_row=8, values_only=True)
header_row = next(rows_iter)
headers = {str(v).strip().lower(): i for i, v in enumerate(header_row) if v}
print(f"Headers: {headers}")

tipo_idx = headers.get('tipo', 9)
valor_idx = headers.get('valor', 11)
cpf_idx = headers.get('cpf', 12)
cod_idx = headers.get('código de transação', 5)
desc_idx = headers.get('descrição', 10)
data_idx = headers.get('data', 3)

# The codes I claimed were "the same" as Saque entries
suspect_codes = {
    '5BE549C9FE239179BFF0EAF8CF470C55',
    '1271F9ADA44C34BF9B9F74A8973F358D',
    '77D17B1E01FCD933209BE915223F9191',
    '8B3268CBE4B4C60F994B171EB803C3BC',
    '3F9DECB7C47979C995ED1749359F294E',
    '64400917DB3A97DB37C1B7448744D9DF',
}

# Collect ALL entries from planilha that have these codes
print("\n=== Planilha entries with suspect codes ===")
for row in rows_iter:
    if row is None:
        continue
    cod_val = str(row[cod_idx] or '').strip() if cod_idx < len(row) else ''
    if cod_val in suspect_codes:
        cpf_val = str(row[cpf_idx] or '').strip() if cpf_idx < len(row) else ''
        tipo_val = str(row[tipo_idx] or '').strip() if tipo_idx < len(row) else ''
        valor_val = float(row[valor_idx] or 0) if valor_idx < len(row) else 0.0
        data_val = str(row[data_idx] or '') if data_idx < len(row) else ''
        desc_val = str(row[desc_idx] or '') if desc_idx < len(row) else ''
        print(f"  CPF={cpf_val}, Data={data_val}, Tipo={tipo_val}, Valor={valor_val}, Cod={cod_val}, Desc={desc_val}")

# Now find the "Taxa pendente de SAQUE" entries and their codes
print("\n=== All 'Taxa pendente de SAQUE' entries in planilha ===")
ws2 = wb['EXTRATO']
rows2 = ws2.iter_rows(min_row=9, values_only=True)
taxa_pendente_codes = []
for row in rows2:
    if row is None:
        continue
    desc_val = str(row[desc_idx] or '').strip() if desc_idx < len(row) else ''
    if 'taxa pendente' in desc_val.lower() and 'saque' in desc_val.lower():
        cpf_val = str(row[cpf_idx] or '').strip() if cpf_idx < len(row) else ''
        tipo_val = str(row[tipo_idx] or '').strip() if tipo_idx < len(row) else ''
        valor_val = float(row[valor_idx] or 0) if valor_idx < len(row) else 0.0
        data_val = str(row[data_idx] or '') if data_idx < len(row) else ''
        cod_val = str(row[cod_idx] or '').strip() if cod_idx < len(row) else ''
        print(f"  CPF={cpf_val}, Data={data_val}, Tipo={tipo_val}, Valor={valor_val}, Cod={cod_val}, Desc={desc_val}")
        taxa_pendente_codes.append(cod_val)

# Now for each Taxa pendente code, find ALL entries in planilha with that same code
print("\n=== For each Taxa pendente code, ALL planilha entries with that code ===")
ws3 = wb['EXTRATO']
rows3 = ws3.iter_rows(min_row=9, values_only=True)
all_by_code = {}
for row in rows3:
    if row is None:
        continue
    cod_val = str(row[cod_idx] or '').strip() if cod_idx < len(row) else ''
    if cod_val in taxa_pendente_codes:
        if cod_val not in all_by_code:
            all_by_code[cod_val] = []
        cpf_val = str(row[cpf_idx] or '').strip() if cpf_idx < len(row) else ''
        tipo_val = str(row[tipo_idx] or '').strip() if tipo_idx < len(row) else ''
        valor_val = float(row[valor_idx] or 0) if valor_idx < len(row) else 0.0
        data_val = str(row[data_idx] or '') if data_idx < len(row) else ''
        desc_val = str(row[desc_idx] or '') if desc_idx < len(row) else ''
        all_by_code[cod_val].append((cpf_val, data_val, tipo_val, valor_val, desc_val))

for cod, entries in all_by_code.items():
    print(f"\n  Code: {cod}")
    for e in entries:
        print(f"    CPF={e[0]}, Data={e[1]}, Tipo={e[2]}, Valor={e[3]}, Desc={e[4]}")
