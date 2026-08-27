import openpyxl
import psycopg2
import warnings
warnings.filterwarnings('ignore')

PLANILHA = 'CONTROLE - VEXPENSES - AGOSTO 2026.xlsx'
COMPARACAO = 'COMPARACAO_BANCO_PLANILHA_v3.xlsx'
DB_CONN = "postgresql://neondb_owner:npg_iItZN95svyEG@ep-restless-voice-amrrz188-pooler.c-5.us-east-1.aws.neon.tech/neondb?sslmode=require"

# Step 1: Get all tarifa diffs from comparison sheet
print("=== Reading comparison sheet ===")
wb = openpyxl.load_workbook(COMPARACAO, read_only=True, data_only=True)
ws = wb['Diff_PAINEL']

tarifa_diffs = []
rows = ws.iter_rows(min_row=5, values_only=True)  # data starts at row 5
for row in rows:
    if row is None or row[0] is None:
        continue
    try:
        tarifa_diff = float(row[10]) if row[10] is not None else 0
    except (ValueError, TypeError):
        continue
    if abs(tarifa_diff) > 0.01:
        tarifa_diffs.append({
            'cpf': str(row[0]).strip(),
            'nome': str(row[1]).strip(),
            'tarifa_plan': float(row[8]),
            'tarifa_banco': float(row[9]),
            'diff': tarifa_diff,
        })

print(f"Found {len(tarifa_diffs)} tarifa diffs")

# Step 2: For each diff, get planilha TARIFA entries and banco Taxa+estorno entries
print("\n=== Reading planilha TARIFA entries ===")
wb2 = openpyxl.load_workbook(PLANILHA, read_only=True, data_only=True)
ws_ext = wb2['EXTRATO']

rows_iter = ws_ext.iter_rows(min_row=8, values_only=True)
header_row = next(rows_iter)
headers = {str(v).strip().lower(): i for i, v in enumerate(header_row) if v}

tipo_idx = headers.get('tipo', 9)
valor_idx = headers.get('valor', 11)
cpf_idx = headers.get('cpf', 12)
cod_idx = headers.get('código de transação', 5)
desc_idx = headers.get('descrição', 10)
data_idx = headers.get('data', 3)

# Collect ALL TARIFA entries from planilha, keyed by CPF
planilha_tarifa = {}  # cpf -> {codigo: (data, valor, desc)}
for row in rows_iter:
    if row is None:
        continue
    cpf_val = str(row[cpf_idx] or '').strip() if cpf_idx < len(row) else ''
    tipo_val = str(row[tipo_idx] or '').strip() if tipo_idx < len(row) else ''
    if tipo_val != 'TARIFA':
        continue
    cod_val = str(row[cod_idx] or '').strip() if cod_idx < len(row) else ''
    valor_val = float(row[valor_idx] or 0) if valor_idx < len(row) else 0.0
    data_val = str(row[data_idx] or '') if data_idx < len(row) else ''
    desc_val = str(row[desc_idx] or '') if desc_idx < len(row) else ''
    
    if cpf_val not in planilha_tarifa:
        planilha_tarifa[cpf_val] = {}
    planilha_tarifa[cpf_val][cod_val] = (data_val, valor_val, desc_val)

# Step 3: Connect to banco and get Taxa + estorno entries for each diff CPF
print("\n=== Connecting to banco ===")
conn = psycopg2.connect(DB_CONN)
cur = conn.cursor()

diff_cpfs = {d['cpf'] for d in tarifa_diffs}

# Get banco entries (Taxa + estorno-as-Transferencia) for these CPFs
# Need to map CPF to usuario name via cadastro
cur.execute("""
    SELECT cpf, colaborador FROM quinzena_cadastro
""")
cpf_to_name = {}
name_to_cpf = {}
for cpf, nome in cur.fetchall():
    cpf = cpf.strip()
    cpf_to_name[cpf] = nome.strip() if nome else ''
    name_to_cpf[nome.strip().upper() if nome else ''] = cpf

# For each diff CPF, get banco Taxa + estorno entries
banco_tarifa = {}  # cpf -> {codigo: (data, valor, tipo, desc)}
for d in tarifa_diffs:
    cpf = d['cpf']
    nome = cpf_to_name.get(cpf, '')
    if not nome:
        print(f"  WARNING: No cadastro for {cpf}")
        continue
    
    cur.execute("""
        SELECT data, tipo, valor, codigo_transacao, descricao
        FROM extrato_movimentacao
        WHERE is_snapshot = false AND data <= '2026-07-31'
        AND UPPER(usuario) = %s
        AND (tipo = 'Taxa' 
             OR tipo = 'Estorno de taxa'
             OR tipo = 'Pendência de taxa'
             OR (tipo = 'Transferência' AND descricao ~* 'estorno.*taxa|taxa.*estorno|^CHARGEBACK_'))
        ORDER BY data
    """, (nome.upper(),))
    
    banco_tarifa[cpf] = {}
    for data, tipo, valor, cod, desc in cur.fetchall():
        banco_tarifa[cpf][cod or ''] = (str(data), float(valor), tipo, desc or '')

# Step 4: Compare and classify diffs
print("\n=== Analysis Results ===\n")
estorno_related_count = 0
api_missing_count = 0
planilha_missing_count = 0

for d in tarifa_diffs:
    cpf = d['cpf']
    nome = d['nome']
    diff = d['diff']
    
    plan_entries = planilha_tarifa.get(cpf, {})
    banco_entries = banco_tarifa.get(cpf, {})
    
    plan_codes = set(plan_entries.keys())
    banco_codes = set(banco_entries.keys())
    
    missing_from_banco = plan_codes - banco_codes  # in planilha but not banco
    extra_in_banco = banco_codes - plan_codes       # in banco but not planilha
    
    # Classify missing from banco
    missing_estorno = 0
    missing_taxa = 0
    missing_value = 0
    for cod in missing_from_banco:
        _, val, desc = plan_entries[cod]
        missing_value += val
        if 'estorno' in desc.lower() or 'chargeback' in desc.lower() or val > 0:
            missing_estorno += 1
        else:
            missing_taxa += 1
    
    # Classify extra in banco
    extra_estorno = 0
    extra_taxa = 0
    extra_value = 0
    for cod in extra_in_banco:
        _, val, tipo, desc = banco_entries[cod]
        extra_value += val
        if 'estorno' in desc.lower() or 'chargeback' in desc.lower() or val > 0:
            extra_estorno += 1
        else:
            extra_taxa += 1
    
    is_estorno_related = (missing_estorno > 0 or extra_estorno > 0)
    if is_estorno_related:
        estorno_related_count += 1
    
    # Check if the missing entries are returned by the API
    # (they're in the banco but as different type, or they're truly missing)
    api_returns_missing = len(extra_in_banco) > 0  # API returned entries not in planilha
    
    print(f"{cpf} {nome}:")
    print(f"  Planilha tarifa: {d['tarifa_plan']}, Banco tarifa: {d['tarifa_banco']}, Diff: {diff}")
    print(f"  Planilha entries: {len(plan_codes)}, Banco entries: {len(banco_codes)}")
    print(f"  Missing from banco: {len(missing_from_banco)} ({missing_taxa} taxa, {missing_estorno} estorno, value={missing_value})")
    if missing_from_banco:
        for cod in sorted(missing_from_banco):
            _, val, desc = plan_entries[cod]
            etype = 'ESTORNO' if val > 0 or 'estorno' in desc.lower() or 'chargeback' in desc.lower() else 'TAXA'
            print(f"    [{etype}] {cod}: val={val}, desc={desc[:50]}")
    print(f"  Extra in banco: {len(extra_in_banco)} ({extra_taxa} taxa, {extra_estorno} estorno, value={extra_value})")
    if extra_in_banco:
        for cod in sorted(extra_in_banco):
            _, val, tipo, desc = banco_entries[cod]
            etype = 'ESTORNO' if val > 0 or 'estorno' in desc.lower() or 'chargeback' in desc.lower() else 'TAXA'
            print(f"    [{etype}] {cod}: val={val}, tipo={tipo}, desc={desc[:50]}")
    print(f"  Estorno-related: {is_estorno_related}")
    print()

conn.close()

print(f"\n=== Summary ===")
print(f"Total tarifa diffs: {len(tarifa_diffs)}")
print(f"Estorno-related: {estorno_related_count}")
print(f"Not estorno-related (pure data gap): {len(tarifa_diffs) - estorno_related_count}")
