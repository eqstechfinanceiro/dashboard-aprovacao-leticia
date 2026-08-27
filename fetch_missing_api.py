import requests
import json

API_URL = 'https://api.vexpenses.com'

# Token from vexpenses_tokens table
LARAVEL_TOKEN = "eyJpdiI6IlBXOC9qcXpaQmlQTDBIUzJPWE95REE9PSIsInZhbHVlIjoiVUM5MTArSGQ1UXd5Y0xrbjQ2U3RhV2dONmtEZ2NzTVE1dmF2TEdMWkx5YitSMWszWkNGdk9oTHVnNW1xMkZhMVRHVTRjU2R5blBtUjVTNE1tK0JQTXByd2FIYmRCZmR3cS9pN01SRWVORjVldlhSUFVZNHpyZDZGLzJVbGk1QjBSano2anVqa0xzY1JiMVBCUHZKaHlJOTR5ek5FSFRRQkQ4NUJiWW5qRGxVSE4vb3dWN2pKaW9qbkJyM3NkY1VOZS9WV0d2TnFBQ2o3QlBudENkUUMzVy94SjcxWTFwVlFHdUIwNm5aZHZJWTdYWi9IYWU4cWxlZWhOT21HUTNGL0JVM2EyTDNod1dyN1pwd1dyUmVJc09SV01VNldEazBRellHOFdXREdvcDIvMHNRbDhMOEJWbzAzVWpLLzkxaXQiLCJtYWMiOiI3ODRhNmQwZmMzMjA3ZDkxOTM3N2NmOWRhNThlOTJjNmNhMzY4MjU4OWU4YzU3YjkwMTAzZmQ1Mjk0OWVjNWExIiwidGFnIjoiIn0%3D"
LARAVEL_SESSION = "m25jkuQuhvM3PsG8ENnYTauKMd8kbHmDIeyLSRqv"
XSRF_TOKEN = "eyJpdiI6IkFlYm82amw0Q3p0T2l5QnMzV0tRU1E9PSIsInZhbHVlIjoiZ09jLzYwWXdjaFdreXdjQ3BERXFvVTl2azRldFJaaUxvOFNHbENBYkVwbUpYYk05ZzFlc3FGRDAzeHphTEpSL3ZvdlFwRG83ajBKZ2NyWE44cW1xWFJaWXdudUJ1ZG90VTBmVWtWbVlrRGpCNXZpakxoYVBENXRYRU5MeC9rMzkiLCJtYWMiOiI2ZDRjMWJhYTIwNjM2ODBmYzRkNWQ3YWUwYTU2OTMxMzEyZmY2ZjE3Yjk4ZmY5MzUzNGI2ZDhkMTZkNTYzMGE3IiwidGFnIjoiIn0%3D"

COOKIE = f"laravel_token={LARAVEL_TOKEN}; laravel_session={LARAVEL_SESSION}; XSRF-TOKEN={XSRF_TOKEN}; language=pt-BR"

HEADERS = {
    'Cookie': COOKIE,
    'Accept': 'application/json',
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
    'Origin': 'https://app.vexpenses.com',
    'Referer': 'https://app.vexpenses.com',
}

# Missing entries to find:
# JORGE: 2026-04-25, code 5AF49CCF32FBA587E6B3A866FA862461
# FABIO: 2026-07-24, code 88B7FCA3CBC6281F441E6DC65179037E
# FERNANDO: 2025-10-27, code D0F2037A63CFDB83E8007AEC1C367DB4
# FELIPE: 2026-06-09, code 2C1829A64681002E5636D17C4C85D09C
# IVON: 2026-03-11, code D2F0E6F3D6C3D94E3299BED09F8A4451

missing = [
    ('JORGE', '2026-04-25', '5AF49CCF32FBA587E6B3A866FA862461'),
    ('FABIO', '2026-07-24', '88B7FCA3CBC6281F441E6DC65179037E'),
    ('FERNANDO', '2025-10-27', 'D0F2037A63CFDB83E8007AEC1C367DB4'),
    ('FELIPE', '2026-06-09', '2C1829A64681002E5636D17C4C85D09C'),
    ('IVON', '2026-03-11', 'D2F0E6F3D6C3D94E3299BED09F8A4451'),
]

# Fetch in 15-day chunks covering each date
chunks = [
    ('2025-10-20', '2025-11-03', 'FERNANDO'),
    ('2026-03-05', '2026-03-19', 'IVON'),
    ('2026-04-20', '2026-05-04', 'JORGE'),
    ('2026-06-04', '2026-06-18', 'FELIPE'),
    ('2026-07-20', '2026-08-03', 'FABIO'),
]

import openpyxl
import io
import warnings
warnings.filterwarnings('ignore')

for chunk_start, chunk_end, label in chunks:
    print(f"\n=== Fetching {label}: {chunk_start} to {chunk_end} ===")
    
    # Step 1: Get S3 presigned URL
    url = f"{API_URL}/v3/pay/statement/excel-all?start_date={chunk_start}&end_date={chunk_end}"
    print(f"  Requesting: {url}")
    
    try:
        resp = requests.get(url, headers=HEADERS, timeout=90)
        print(f"  Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"  Error: {resp.text[:500]}")
            continue
        
        data = resp.json()
        if not data.get('success'):
            print(f"  API error: {json.dumps(data)[:500]}")
            continue
        
        s3_url = data.get('data', {}).get('url')
        if not s3_url:
            print(f"  No S3 URL in response: {json.dumps(data)[:500]}")
            continue
        
        print(f"  S3 URL obtained, downloading XLSX...")
        
        # Step 2: Download XLSX
        xlsx_resp = requests.get(s3_url, timeout=180)
        print(f"  XLSX download status: {xlsx_resp.status_code}")
        
        if xlsx_resp.status_code != 200:
            print(f"  XLSX download failed")
            continue
        
        # Step 3: Parse XLSX
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        
        rows = ws.iter_rows(min_row=1, values_only=True)
        header = next(rows)
        print(f"  Headers: {header}")
        
        # Find column indices
        header_map = {str(v).strip().lower(): i for i, v in enumerate(header) if v}
        cod_idx = header_map.get('código de transação', 4)
        tipo_idx = header_map.get('tipo', 7)
        desc_idx = header_map.get('descrição', 8)
        valor_idx = header_map.get('valor', 9)
        user_idx = header_map.get('usuário', 6)
        
        # Search for our missing codes
        found_codes = set()
        row_count = 0
        taxa_count = 0
        for row in rows:
            if row is None:
                continue
            row_count += 1
            cod_val = str(row[cod_idx] or '').strip() if cod_idx < len(row) else ''
            tipo_val = str(row[tipo_idx] or '').strip() if tipo_idx < len(row) else ''
            
            if tipo_val == 'Taxa':
                taxa_count += 1
            
            # Check if this row has one of our missing codes
            for label_name, _, missing_code in missing:
                if cod_val == missing_code:
                    desc_val = str(row[desc_idx] or '') if desc_idx < len(row) else ''
                    valor_val = row[valor_idx] if valor_idx < len(row) else ''
                    user_val = str(row[user_idx] or '') if user_idx < len(row) else ''
                    print(f"  *** FOUND: {label_name} code={cod_val}")
                    print(f"      Tipo={tipo_val}, Desc={desc_val}, Valor={valor_val}, User={user_val}")
                    found_codes.add(cod_val)
        
        print(f"  Total rows: {row_count}, Taxa entries: {taxa_count}")
        print(f"  Found codes: {found_codes}")
        
        # Check which ones we didn't find
        expected_code = [c for l, d, c in missing if l == label][0]
        if expected_code not in found_codes:
            print(f"  *** NOT FOUND in API response: {expected_code}")
        
    except Exception as e:
        print(f"  Exception: {e}")

print("\n=== Done ===")
